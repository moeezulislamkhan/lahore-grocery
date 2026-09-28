# app/routes/orders.py
import secrets
from datetime import datetime
from io import BytesIO

from flask import Blueprint, request, jsonify, current_app, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.db import query_one, query_all, execute
from app.sessions import get_session_user, role_allowed
from app.payments_jazzcash import build_checkout_form
from app.deal_utils import active_deal_for_product

bp = Blueprint('orders', __name__, url_prefix='/api/orders')

VALID_METHODS = {'jazzcash', 'bank', 'cod'}
VALID_STATUSES = ['processing', 'confirmed', 'out_for_delivery', 'delivered', 'returned']
VALID_PAYMENT_STATUSES = ['pending', 'paid', 'failed', 'cod_pending', 'refunded']


def generate_order_code():
    return 'SG-' + datetime.now().strftime('%y%m%d%H%M%S') + '-' + secrets.token_hex(2).upper()


@bp.post('')
def create_order():
    db = current_app.get_db()
    body = request.get_json(silent=True) or {}
    customer_name, phone, address = body.get('customerName'), body.get('phone'), body.get('address')
    city, delivery_slot = body.get('city'), body.get('deliverySlot')
    items, payment_method = body.get('items'), body.get('paymentMethod')
    delivery_instructions = body.get('deliveryInstructions') or body.get('delivery_instructions')
    postal_code = body.get('postalCode') or body.get('postal_code')
    customer_email = (body.get('email') or '').strip() or ((get_session_user(db, request) or {}).get('email') if get_session_user(db, request) else '')
    customer_phone = body.get('customerPhone') or body.get('phone')

    if not customer_name or not phone or not address:
        return jsonify(error='customerName, phone and address are required.'), 400
    if city and city.strip().lower() != 'lahore':
        return jsonify(error='We currently deliver only within Lahore.'), 400
    city = 'Lahore'
    if not items or not isinstance(items, list):
        return jsonify(error='Cart is empty.'), 400
    if payment_method not in VALID_METHODS:
        return jsonify(error='Invalid payment method.'), 400

    session_user = get_session_user(db, request)  # None for guest checkout — that's fine

    subtotal = 0.0
    resolved_items = []
    for item in items:
        product = query_one(db, 'SELECT * FROM products WHERE id = ?', (item.get('productId'),))
        if not product:
            return jsonify(error=f"Product {item.get('productId')} not found."), 400
        qty = max(1, int(item.get('qty') or 1))
        deal = active_deal_for_product(db, product['id'])
        unit_price = deal['final_price'] if deal else float(product['sale_price'] or product['price'])
        subtotal += unit_price * qty
        resolved_items.append((product, qty, item.get('variant') or product.get('unit') or 'Standard'))

    delivery_fee = 0.0 if subtotal > 2000 else 150.0
    total = subtotal + delivery_fee
    order_code = generate_order_code()
    payment_status = 'cod_pending' if payment_method == 'cod' else 'pending'

    cur = execute(db, """INSERT INTO orders
        (order_code, user_id, customer_name, customer_email, customer_phone, phone, address, city, postal_code, delivery_slot, delivery_instructions, subtotal, delivery_fee, total, payment_method, payment_status)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (order_code, session_user['id'] if session_user else None, customer_name, customer_email or None, customer_phone or phone, phone, address, city,
         postal_code, delivery_slot, delivery_instructions, subtotal, delivery_fee, total, payment_method, payment_status))
    order_id = cur.lastrowid

    for product, qty, variant in resolved_items:
        deal = active_deal_for_product(db, product['id'])
        unit_price = deal['final_price'] if deal else float(product['sale_price'] or product['price'])
        line_total = unit_price * qty
        execute(db, 'INSERT INTO order_items (order_id, product_id, product_name, product_name_snapshot, product_image_snapshot, variant_snapshot, unit_price, qty, discount, line_total) VALUES (?,?,?,?,?,?,?,?,?,?)',
                (order_id, product['id'], product['name'], product['name'], product.get('image'), variant, unit_price, qty, 0.0, line_total))
        execute(db, 'UPDATE products SET stock = MAX(0, stock - ?) WHERE id = ?', (qty, product['id']))

    if payment_method == 'cod':
        if hasattr(db, 'commit'):
            db.commit()
        return jsonify(orderCode=order_code, paymentMethod=payment_method, paymentStatus='cod_pending',
                        message='Order placed — pay cash on delivery.'), 201

    if payment_method == 'bank':
        if hasattr(db, 'commit'):
            db.commit()
        store_name = query_one(db, "SELECT value FROM settings WHERE setting_key = ?", ('store_name',))
        bank_name = query_one(db, "SELECT value FROM settings WHERE setting_key = ?", ('bank_name',))
        bank_iban = query_one(db, "SELECT value FROM settings WHERE setting_key = ?", ('bank_iban',))
        return jsonify(orderCode=order_code, paymentMethod=payment_method, paymentStatus='pending',
                        bankInstructions={
                            'accountTitle': store_name['value'] if store_name else 'Shakarganj Grocery Store',
                            'bank': bank_name['value'] if bank_name else 'Meezan Bank',
                            'iban': bank_iban['value'] if bank_iban else 'PK00 MEZN 0000 0000 1234 567',
                        }), 201

    checkout = build_checkout_form(
        order_code=order_code,
        amount_pkr=total,
        bill_reference=order_code,
        description=f'Shakarganj order {order_code}',
        base_url=request.host_url.rstrip('/'),
    )
    execute(db, 'UPDATE orders SET gateway = ?, gateway_txn_ref = ? WHERE id = ?', ('jazzcash', checkout['txn_ref_no'], order_id))
    execute(db, 'INSERT INTO jazzcash_transactions (txn_ref_no, order_id) VALUES (?, ?)', (checkout['txn_ref_no'], order_id))
    if hasattr(db, 'commit'):
        db.commit()

    return jsonify(
        orderCode=order_code,
        paymentMethod=payment_method,
        paymentStatus='pending',
        gatewayLive=checkout['live'],
        amountPkr=round(total, 2),
        jazzcash={'actionUrl': checkout['action_url'], 'fields': checkout['fields']},
    ), 201


@bp.get('')
def list_all_orders():
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Staff access required.'), 403
    orders = query_all(db, 'SELECT * FROM orders ORDER BY created_at DESC LIMIT 200')
    return jsonify(orders=orders)


def _get_order_history(db):
    orders = query_all(db, """SELECT o.*,
        COALESCE(NULLIF(o.customer_email, ''), u.email, '') AS history_customer_email,
        COALESCE(NULLIF(o.customer_phone, ''), o.phone, '') AS history_customer_phone
        FROM orders o LEFT JOIN users u ON u.id = o.user_id
        ORDER BY o.created_at DESC, o.id DESC""")
    items = query_all(db, """SELECT oi.*, o.order_code, o.created_at, o.customer_name
        FROM order_items oi JOIN orders o ON o.id = oi.order_id
        ORDER BY o.created_at DESC, o.id DESC, oi.id ASC""")
    items_by_order = {}
    for item in items:
        items_by_order.setdefault(item['order_id'], []).append(item)

    for order in orders:
        order['customer_email'] = order.pop('history_customer_email')
        order['customer_phone'] = order.pop('history_customer_phone')
        order['items'] = items_by_order.get(order['id'], [])
        order['item_count'] = len(order['items'])
        order['total_quantity'] = sum(int(item.get('qty') or 0) for item in order['items'])
    return orders


def _style_history_sheet(sheet, widths):
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='620014')
        cell.font = Font(color='FFFFFF', bold=True)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _spreadsheet_text(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


@bp.get('/history')
def get_order_history():
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Staff access required.'), 403
    orders = _get_order_history(db)
    return jsonify(orders=orders, order_count=len(orders))


@bp.get('/history.xlsx')
def download_order_history():
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Staff access required.'), 403

    orders = _get_order_history(db)
    workbook = Workbook()
    orders_sheet = workbook.active
    orders_sheet.title = 'Orders'
    orders_sheet.append([
        'Order Code', 'Placed At', 'Customer Name', 'Customer Email', 'Phone',
        'Address', 'City', 'Postal Code', 'Delivery Slot', 'Delivery Instructions',
        'Products', 'Product Lines', 'Total Quantity', 'Subtotal', 'Delivery Fee',
        'Total', 'Payment Method', 'Payment Status', 'Order Status', 'Gateway',
        'Transaction Reference',
    ])
    items_sheet = workbook.create_sheet('Order Items')
    items_sheet.append([
        'Order Code', 'Placed At', 'Customer Name', 'Product ID', 'Product Name',
        'Variant', 'Quantity', 'Unit Price', 'Discount', 'Line Total',
    ])

    for order in orders:
        products = '; '.join(
            f"{item.get('product_name_snapshot') or item.get('product_name') or 'Product'} x{item.get('qty') or 0}"
            for item in order['items']
        )
        orders_sheet.append([
            *[_spreadsheet_text(order.get(key) or '') for key in (
                'order_code', 'created_at', 'customer_name', 'customer_email',
                'customer_phone', 'address', 'city', 'postal_code',
                'delivery_slot', 'delivery_instructions',
            )],
            _spreadsheet_text(products), order['item_count'], order['total_quantity'],
            order.get('subtotal') or 0, order.get('delivery_fee') or 0,
            order.get('total') or 0, _spreadsheet_text(order.get('payment_method') or ''),
            _spreadsheet_text(order.get('payment_status') or ''),
            _spreadsheet_text(order.get('order_status') or ''),
            _spreadsheet_text(order.get('gateway') or ''),
            _spreadsheet_text(order.get('gateway_txn_ref') or ''),
        ])
        for item in order['items']:
            items_sheet.append([
                _spreadsheet_text(order.get('order_code') or ''),
                _spreadsheet_text(order.get('created_at') or ''),
                _spreadsheet_text(order.get('customer_name') or ''),
                item.get('product_id') or '',
                _spreadsheet_text(item.get('product_name_snapshot') or item.get('product_name') or ''),
                _spreadsheet_text(item.get('variant_snapshot') or ''),
                item.get('qty') or 0, item.get('unit_price') or 0,
                item.get('discount') or 0, item.get('line_total') or 0,
            ])

    _style_history_sheet(orders_sheet, {
        'A': 24, 'B': 21, 'C': 25, 'D': 32, 'E': 18, 'F': 38, 'G': 16,
        'H': 16, 'I': 20, 'J': 32, 'K': 48, 'L': 14, 'M': 15, 'N': 15,
        'O': 15, 'P': 15, 'Q': 18, 'R': 18, 'S': 20, 'T': 16, 'U': 25,
    })
    _style_history_sheet(items_sheet, {
        'A': 24, 'B': 21, 'C': 25, 'D': 14, 'E': 38,
        'F': 20, 'G': 12, 'H': 15, 'I': 13, 'J': 15,
    })
    for sheet, money_columns in ((orders_sheet, ('N', 'O', 'P')), (items_sheet, ('H', 'I', 'J'))):
        for column in money_columns:
            for cell in sheet[column][1:]:
                cell.number_format = '#,##0.00'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"shakarganj_order_history_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@bp.get('/mine')
def list_my_orders():
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not user:
        return jsonify(error='Please sign in to view your orders.'), 401
    orders = query_all(db, 'SELECT * FROM orders WHERE user_id = ? ORDER BY created_at DESC LIMIT 100', (user['id'],))
    return jsonify(orders=orders)


@bp.patch('/<int:order_id>/cancel')
def cancel_my_order(order_id):
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not user:
        return jsonify(error='Please sign in to cancel this order.'), 401
    if user['role'] != 'user':
        return jsonify(error='Customer access required.'), 403

    order = query_one(db, 'SELECT * FROM orders WHERE id = ? AND user_id = ?', (order_id, user['id']))
    if not order:
        return jsonify(error='Order not found.'), 404
    if order['order_status'] == 'cancelled':
        return jsonify(error='This order has already been cancelled.'), 409
    if order['order_status'] in ['out_for_delivery', 'delivered', 'returned']:
        return jsonify(error="You're unable to cancel this order because this order is on the way"), 409

    result = execute(db, """UPDATE orders SET order_status = 'cancelled'
        WHERE id = ? AND user_id = ? AND order_status IN ('processing', 'confirmed')""",
        (order_id, user['id']))
    if result.rowcount == 0:
        current_order = query_one(db, 'SELECT order_status FROM orders WHERE id = ? AND user_id = ?', (order_id, user['id']))
        if current_order and current_order['order_status'] in ['out_for_delivery', 'delivered', 'returned']:
            return jsonify(error="You're unable to cancel this order because this order is on the way"), 409
        if current_order and current_order['order_status'] == 'cancelled':
            return jsonify(error='This order has already been cancelled.'), 409
        return jsonify(error='This order cannot be cancelled at its current status.'), 409

    if hasattr(db, 'commit'):
        db.commit()
    updated_order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    return jsonify(success=True, order=updated_order)


def _serialize_order_detail(db, order):
    items = query_all(db, 'SELECT * FROM order_items WHERE order_id = ?', (order['id'],))
    user = query_one(db, 'SELECT id, name, email FROM users WHERE id = ?', (order['user_id'],)) if order.get('user_id') else None
    detail = dict(order)
    detail['customer_email'] = detail.get('customer_email') or (user['email'] if user else 'Not Provided')
    detail['customer_phone'] = detail.get('customer_phone') or detail.get('phone') or 'Not Provided'
    detail['delivery_address'] = detail.get('address') or 'Not Provided'
    detail['postal_code'] = detail.get('postal_code') or 'Not Provided'
    detail['delivery_instructions'] = detail.get('delivery_instructions') or 'Not Provided'
    detail['payment_method_display'] = detail.get('payment_method', 'cod').replace('_', ' ').title() if detail.get('payment_method') else 'Not Provided'
    detail['payment_status_display'] = (detail.get('payment_status') or 'pending').replace('_', ' ').title()
    detail['order_status_display'] = (detail.get('order_status') or 'processing').replace('_', ' ').title()
    detail['customer'] = {
        'id': user['id'] if user else None,
        'full_name': detail.get('customer_name') or (user['name'] if user else 'Not Provided'),
        'email': detail['customer_email'],
        'phone': detail['customer_phone'],
        'address': detail['delivery_address'],
        'city': detail.get('city') or 'Not Provided',
        'postal_code': detail['postal_code'],
    }
    detail['delivery'] = {
        'address': detail['delivery_address'],
        'city': detail.get('city') or 'Not Provided',
        'postal_code': detail['postal_code'],
        'instructions': detail['delivery_instructions'],
        'slot': detail.get('delivery_slot') or 'Not Provided',
    }
    detail['payment'] = {
        'method': detail.get('payment_method') or 'Not Provided',
        'status': detail.get('payment_status') or 'pending',
        'transaction_reference': detail.get('gateway_txn_ref') or 'Not Provided',
        'gateway': detail.get('gateway') or 'Not Provided',
    }
    return {'order': detail, 'items': items}


@bp.get('/<code>')
def get_order(code):
    db = current_app.get_db()
    order = None
    if str(code).isdigit():
        order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (int(code),))
    if not order:
        order = query_one(db, 'SELECT * FROM orders WHERE order_code = ?', (code,))
    if not order:
        return jsonify(error='Order not found.'), 404
    payload = _serialize_order_detail(db, order)
    return jsonify(order=payload['order'], items=payload['items'])


@bp.get('/<int:order_id>/detail')
def get_order_detail(order_id):
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Admin access required.'), 403
    order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    if not order:
        return jsonify(error='Order not found.'), 404
    payload = _serialize_order_detail(db, order)
    return jsonify(payload)


@bp.get('/<int:order_id>/receipt')
def get_order_receipt(order_id):
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Admin access required.'), 403
    order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    if not order:
        return jsonify(error='Order not found.'), 404
    payload = _serialize_order_detail(db, order)
    return jsonify({'order': payload['order'], 'items': payload['items'], 'receipt': {
        'order_code': payload['order']['order_code'],
        'customer_name': payload['order']['customer']['full_name'],
        'subtotal': payload['order'].get('subtotal') or 0,
        'delivery_fee': payload['order'].get('delivery_fee') or 0,
        'total': payload['order'].get('total') or 0,
    }})


@bp.patch('/<int:order_id>/status')
def update_order_status(order_id):
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Staff access required.'), 403
    body = request.get_json(silent=True) or {}
    status = body.get('status')
    if status not in VALID_STATUSES:
        return jsonify(error='Invalid status.'), 400
    current_order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    if not current_order:
        return jsonify(error='Order not found.'), 404
    execute(db, 'UPDATE orders SET order_status = ? WHERE id = ?', (status, order_id))
    if hasattr(db, 'commit'):
        db.commit()
    order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    return jsonify(success=True, order=order)


@bp.patch('/<int:order_id>/payment-status')
def update_payment_status(order_id):
    db = current_app.get_db()
    user = get_session_user(db, request)
    if not role_allowed(user, ['admin', 'manager', 'staff', 'employee']):
        return jsonify(error='Admin access required.'), 403
    body = request.get_json(silent=True) or {}
    payment_status = body.get('payment_status')
    if payment_status not in VALID_PAYMENT_STATUSES:
        return jsonify(error='Invalid payment status.'), 400
    order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    if not order:
        return jsonify(error='Order not found.'), 404
    transaction_reference = body.get('transaction_reference') or order.get('gateway_txn_ref')
    execute(db, 'UPDATE orders SET payment_status = ?, gateway_txn_ref = ? WHERE id = ?', (payment_status, transaction_reference, order_id))
    if hasattr(db, 'commit'):
        db.commit()
    order = query_one(db, 'SELECT * FROM orders WHERE id = ?', (order_id,))
    return jsonify(success=True, order=order)
