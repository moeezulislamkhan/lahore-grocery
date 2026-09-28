from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app import create_app
from app.auth import hash_password
from app.db import get_connection, query_one


def _login_admin(client):
    response = client.post('/api/auth/login', json={'email': 'shakarganj@gmail.com', 'password': '11223344'})
    assert response.status_code == 200, response.get_data(as_text=True)
    return response.get_json()


def _login_customer(client):
    response = client.post('/api/auth/login', json={'email': 'customer@shakarganj.pk', 'password': 'Customer@123'})
    assert response.status_code == 200, response.get_data(as_text=True)
    return response.get_json()


def test_admin_order_detail_and_payment_status_endpoints():
    app = create_app()
    with app.app_context():
        db = app.get_db()
        product = query_one(db, 'SELECT * FROM products ORDER BY id LIMIT 1')
        assert product is not None
        order_code = 'TEST-ORDER-DETAIL-1'
        cur = db.execute(
            '''INSERT INTO orders
               (order_code, user_id, customer_name, phone, address, city, subtotal, delivery_fee, total, payment_method, payment_status, order_status, customer_email, customer_phone, postal_code, delivery_instructions)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
            (order_code, 1, 'Ali Khan', '03001234567', '123 Main St', 'Lahore', 500.0, 150.0, 650.0, 'cod', 'cod_pending', 'confirmed', 'ali@example.com', '03001234567', '54000', 'Leave at gate')
        )
        order_id = cur.lastrowid
        db.execute(
            'INSERT INTO order_items (order_id, product_id, product_name, product_name_snapshot, product_image_snapshot, unit_price, qty, line_total, variant_snapshot) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
            (order_id, product['id'], product['name'], product['name'], product['image'], 250.0, 2, 500.0, '1L')
        )
        db.commit()

    client = app.test_client()
    admin = _login_admin(client)

    detail_response = client.get(f'/api/orders/{order_id}/detail', headers={'Authorization': f"Bearer {admin['token']}"})
    assert detail_response.status_code == 200, detail_response.get_data(as_text=True)
    payload = detail_response.get_json()
    assert payload['order']['order_code'] == order_code
    assert payload['order']['customer_name'] == 'Ali Khan'
    assert payload['items'][0]['product_name'] == product['name']

    status_response = client.patch(f'/api/orders/{order_id}/status', json={'status': 'out_for_delivery'}, headers={'Authorization': f"Bearer {admin['token']}"})
    assert status_response.status_code == 200, status_response.get_data(as_text=True)

    payment_response = client.patch(f'/api/orders/{order_id}/payment-status', json={'payment_status': 'paid'}, headers={'Authorization': f"Bearer {admin['token']}"})
    assert payment_response.status_code == 200, payment_response.get_data(as_text=True)
    assert payment_response.get_json()['order']['payment_status'] == 'paid'

    returned_response = client.patch(f'/api/orders/{order_id}/status', json={'status': 'returned'}, headers={'Authorization': f"Bearer {admin['token']}"})
    assert returned_response.status_code == 200, returned_response.get_data(as_text=True)
    assert returned_response.get_json()['order']['order_status'] == 'returned'

    cancelled_response = client.patch(f'/api/orders/{order_id}/status', json={'status': 'cancelled'}, headers={'Authorization': f"Bearer {admin['token']}"})
    assert cancelled_response.status_code == 400
    assert cancelled_response.get_json()['error'] == 'Invalid status.'


def test_customer_can_cancel_before_dispatch_but_not_after():
    app = create_app()
    with app.app_context():
        db = app.get_db()
        customer = query_one(db, 'SELECT id FROM users WHERE email = ?', ('customer@shakarganj.pk',))
        assert customer is not None
        codes = ['TEST-CANCEL-BEFORE-DISPATCH', 'TEST-CANCEL-AFTER-DISPATCH']
        order_ids = []
        for code, status in zip(codes, ['confirmed', 'out_for_delivery']):
            cur = db.execute(
                '''INSERT INTO orders
                   (order_code, user_id, customer_name, phone, address, city, subtotal, delivery_fee, total, payment_method, payment_status, order_status)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                (code, customer['id'], 'Test Customer', '03001234567', 'Test address', 'Lahore', 500.0, 150.0, 650.0, 'cod', 'cod_pending', status)
            )
            order_ids.append(cur.lastrowid)
        db.commit()

    client = app.test_client()
    customer_session = _login_customer(client)
    headers = {'Authorization': f"Bearer {customer_session['token']}"}
    try:
        cancel_response = client.patch(f'/api/orders/{order_ids[0]}/cancel', headers=headers)
        assert cancel_response.status_code == 200, cancel_response.get_data(as_text=True)
        assert cancel_response.get_json()['order']['order_status'] == 'cancelled'

        blocked_response = client.patch(f'/api/orders/{order_ids[1]}/cancel', headers=headers)
        assert blocked_response.status_code == 409
        assert blocked_response.get_json()['error'] == "You're unable to cancel this order because this order is on the way"
    finally:
        client.post('/api/auth/logout', headers=headers)
        with app.app_context():
            db = app.get_db()
            db.execute('DELETE FROM order_items WHERE order_id IN (?, ?)', tuple(order_ids))
            db.execute('DELETE FROM orders WHERE id IN (?, ?)', tuple(order_ids))
            db.commit()


def test_new_orders_stay_processing_until_admin_updates_status():
    app = create_app()
    with app.app_context():
        db = app.get_db()
        product = query_one(db, 'SELECT id, stock FROM products ORDER BY id LIMIT 1')
        assert product is not None
        original_stock = product['stock']

    client = app.test_client()
    customer = _login_customer(client)
    headers = {'Authorization': f"Bearer {customer['token']}"}
    order_codes = []
    try:
        for payment_method in ['cod', 'jazzcash']:
            response = client.post('/api/orders', json={
                'customerName': 'Test Customer',
                'phone': '03001234567',
                'address': 'Test address',
                'city': 'Lahore',
                'items': [{'productId': product['id'], 'qty': 1}],
                'paymentMethod': payment_method,
            }, headers=headers)
            assert response.status_code == 201, response.get_data(as_text=True)
            order_code = response.get_json()['orderCode']
            order_codes.append(order_code)

            if payment_method == 'jazzcash':
                payment_response = client.post('/api/payments/simulate', json={
                    'orderCode': order_code,
                    'outcome': 'succeeded',
                })
                assert payment_response.status_code == 200, payment_response.get_data(as_text=True)

            order_response = client.get(f'/api/orders/{order_code}')
            assert order_response.status_code == 200, order_response.get_data(as_text=True)
            order = order_response.get_json()['order']
            assert order['order_status'] == 'processing'
            if payment_method == 'jazzcash':
                assert order['payment_status'] == 'paid'
    finally:
        client.post('/api/auth/logout', headers=headers)
        with app.app_context():
            db = app.get_db()
            db.execute('DELETE FROM jazzcash_transactions WHERE order_id IN (SELECT id FROM orders WHERE order_code IN (?, ?))', tuple(order_codes))
            db.execute('DELETE FROM order_items WHERE order_id IN (SELECT id FROM orders WHERE order_code IN (?, ?))', tuple(order_codes))
            db.execute('DELETE FROM orders WHERE order_code IN (?, ?)', tuple(order_codes))
            db.execute('UPDATE products SET stock = ? WHERE id = ?', (original_stock, product['id']))
            db.commit()


def test_admin_order_history_export_contains_order_and_item_sheets():
    app = create_app()
    client = app.test_client()
    admin = _login_admin(client)
    headers = {'Authorization': f"Bearer {admin['token']}"}

    history_response = client.get('/api/orders/history', headers=headers)
    assert history_response.status_code == 200, history_response.get_data(as_text=True)
    history = history_response.get_json()

    export_response = client.get('/api/orders/history.xlsx', headers=headers)
    assert export_response.status_code == 200, export_response.get_data(as_text=True)
    assert export_response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    workbook = load_workbook(BytesIO(export_response.data), read_only=True, data_only=True)
    assert workbook.sheetnames == ['Orders', 'Order Items']
    assert workbook['Orders'].max_row == history['order_count'] + 1
    expected_items = sum(len(order['items']) for order in history['orders'])
    assert workbook['Order Items'].max_row == expected_items + 1
    assert next(workbook['Orders'].iter_rows(min_row=1, max_row=1, values_only=True))[0] == 'Order Code'


def test_admin_alerts_include_stock_support_and_deleted_deal_events():
    app = create_app()
    client = app.test_client()
    admin = _login_admin(client)
    headers = {'Authorization': f"Bearer {admin['token']}"}
    baseline_response = client.get('/api/admin/alerts', headers=headers)
    assert baseline_response.status_code == 200, baseline_response.get_data(as_text=True)
    baseline = baseline_response.get_json()['alerts']

    suffix = uuid4().hex[:10]
    with app.app_context():
        db = app.get_db()
        product_cursor = db.execute(
            'INSERT INTO products (name, category, price, stock) VALUES (?, ?, ?, ?)',
            (f'Alert test product {suffix}', 'Alert tests', 1, 0)
        )
        product_id = product_cursor.lastrowid
        message_cursor = db.execute(
            'INSERT INTO support_messages (name, email, topic, message) VALUES (?, ?, ?, ?)',
            ('Alert Test', f'{suffix}@example.test', 'Alert test', 'Test support alert')
        )
        message_id = message_cursor.lastrowid
        campaign_cursor = db.execute(
            '''INSERT INTO deal_campaigns (title, description, start_date, end_date)
               VALUES (?, ?, ?, ?)''',
            (f'Alert test container {suffix}', 'Test deleted-container alert', '2026-01-01 00:00:00', '2027-01-01 00:00:00')
        )
        campaign_id = campaign_cursor.lastrowid
        db.commit()

    try:
        delete_response = client.delete(f'/api/admin/deal-campaigns/{campaign_id}', headers=headers)
        assert delete_response.status_code == 200, delete_response.get_data(as_text=True)

        alerts_response = client.get('/api/admin/alerts', headers=headers)
        assert alerts_response.status_code == 200, alerts_response.get_data(as_text=True)
        alerts = alerts_response.get_json()['alerts']
        assert alerts['out_of_stock_count'] == baseline['out_of_stock_count'] + 1
        assert alerts['open_support_count'] == baseline['open_support_count'] + 1
        assert alerts['deleted_deal_container_count'] == baseline['deleted_deal_container_count'] + 1
    finally:
        client.post('/api/auth/logout', headers=headers)
        with app.app_context():
            db = app.get_db()
            db.execute('DELETE FROM products WHERE id = ?', (product_id,))
            db.execute('DELETE FROM support_messages WHERE id = ?', (message_id,))
            db.execute("DELETE FROM admin_notifications WHERE notification_type = 'deal_container_deleted' AND entity_id = ?", (campaign_id,))
            db.execute('DELETE FROM deal_campaigns WHERE id = ?', (campaign_id,))
            db.commit()


def test_employee_can_manage_store_sections_but_not_restricted_sections():
    app = create_app()
    email = f"member-{uuid4().hex[:10]}@example.test"
    with app.app_context():
        db = app.get_db()
        cur = db.execute('INSERT INTO users (name, email, password_hash, role) VALUES (?, ?, ?, ?)',
                         ('Test Member', email, hash_password('MemberTest@123'), 'employee'))
        employee_id = cur.lastrowid
        db.commit()

    client = app.test_client()
    employee = client.post('/api/auth/login', json={'email': email, 'password': 'MemberTest@123'}).get_json()
    headers = {'Authorization': f"Bearer {employee['token']}"}
    try:
        for path in (
            '/api/orders', '/api/products', '/api/support/messages',
            '/api/admin/collections', '/api/admin/promo-messages',
            '/api/admin/deals', '/api/admin/deal-campaigns',
        ):
            response = client.get(path, headers=headers)
            assert response.status_code == 200, (path, response.get_data(as_text=True))

        for path in ('/api/analytics/summary', '/api/employees', '/api/settings'):
            response = client.get(path, headers=headers)
            assert response.status_code == 403, (path, response.get_data(as_text=True))
    finally:
        client.post('/api/auth/logout', headers=headers)
        with app.app_context():
            db = app.get_db()
            db.execute('DELETE FROM users WHERE id = ?', (employee_id,))
            db.commit()
